#!/usr/bin/env python3
# ============================================================
#  PROCESAR FOTOS DE PRODUCTOS
# ============================================================
#  Toma fotos sueltas, las empareja con su producto usando el Excel
#  de fotos (hoja "Fotos": N° foto / Codigo de barras / Nombre) y las
#  deja listas en img/productos/<codigo>.webp:
#    - fondo blanco (las transparencias se rellenan de blanco)
#    - cuadradas (se centra el producto, sin recortarlo)
#    - maximo 800x800 y en webp (cargan rapido)
#
#  Uso:
#    python3 herramientas/procesar_fotos.py Fotos.xlsx carpeta_o_fotos...
#        -> cada foto se empareja por su nombre: por el codigo de barras
#           (886540006081.jpg) o, si el Excel tiene N° de foto, tambien
#           por ese N° (0001.jpg = N° 0001).
#
#    python3 herramientas/procesar_fotos.py Fotos.xlsx --desde 6 a.jpg b.jpg c.jpg
#        -> fotos sin nombre util: se asignan en orden desde el N° 6
#           (a.jpg = 0006, b.jpg = 0007, ...). Ojo: si falta una, se corren.
#           Requiere que el Excel tenga la columna N° de foto.
#
#  Despues hay que correr actualizar_catalogo.py para que el catalogo
#  las detecte.
#
#  Requiere: pip install openpyxl pillow
# ============================================================

import argparse
import os
import re
import sys

try:
    import openpyxl
    from PIL import Image, ImageOps
except ImportError:
    sys.exit("Faltan librerias. Instalalas con:  pip install openpyxl pillow")

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DESTINO = os.path.join(RAIZ, "img", "productos")
EXT = (".jpg", ".jpeg", ".png", ".webp")
LADO = 800


def leer_lista(path):
    # Acepta la lista completa (hoja "Fotos": N° foto / Codigo / Nombre)
    # o la lista de pendientes (hoja "Pendientes": Codigo / Nombre / Marca
    # / Nombre del archivo / Categoria / Foto). En la de pendientes no hay
    # N° de foto, asi que solo sirve el emparejado por codigo de barras.
    wb = openpyxl.load_workbook(path, read_only=True)
    nombre_hoja = "Fotos" if "Fotos" in wb.sheetnames else "Pendientes"
    ws = wb[nombre_hoja]
    encabezado = [str(c or "").strip().lower() for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    col_num = encabezado.index("n° foto") if "n° foto" in encabezado else None
    col_code = encabezado.index("código de barras")
    col_nombre = encabezado.index("nombre")

    por_numero, por_codigo = {}, {}
    for fila in ws.iter_rows(min_row=2, values_only=True):
        code = str(fila[col_code] or "").strip()
        if not code:
            continue
        nombre = str(fila[col_nombre] or "").strip()
        por_codigo[code] = (code, nombre)
        if col_num is not None and fila[col_num]:
            por_numero[int(fila[col_num])] = (code, nombre)
    return por_numero, por_codigo


def juntar_fotos(entradas):
    fotos = []
    for e in entradas:
        if os.path.isdir(e):
            fotos += sorted(os.path.join(e, f) for f in os.listdir(e) if f.lower().endswith(EXT))
        else:
            fotos.append(e)
    return fotos


def procesar(src, dst):
    im = ImageOps.exif_transpose(Image.open(src))
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA")
        fondo = Image.new("RGB", im.size, "white")
        fondo.paste(im, mask=im.split()[-1])
        im = fondo
    else:
        im = im.convert("RGB")
    lado = max(im.size)
    cuadro = Image.new("RGB", (lado, lado), "white")
    cuadro.paste(im, ((lado - im.width) // 2, (lado - im.height) // 2))
    if lado > LADO:
        cuadro = cuadro.resize((LADO, LADO), Image.LANCZOS)
    cuadro.save(dst, "WEBP", quality=82)


def main():
    ap = argparse.ArgumentParser(description="Empareja y prepara fotos de productos.")
    ap.add_argument("lista", help="Excel de fotos (hoja 'Fotos')")
    ap.add_argument("fotos", nargs="+", help="Fotos o carpetas con fotos")
    ap.add_argument("--desde", type=int, help="N° de la primera foto si no vienen numeradas")
    args = ap.parse_args()

    por_numero, por_codigo = leer_lista(args.lista)
    os.makedirs(DESTINO, exist_ok=True)
    siguiente = args.desde
    for src in juntar_fotos(args.fotos):
        base = os.path.splitext(os.path.basename(src))[0].strip()
        etiqueta = base
        if base in por_codigo:
            code, nombre = por_codigo[base]
        elif siguiente is not None:
            if siguiente not in por_numero:
                print(f"SALTADA {src}: el N° {siguiente:04d} no esta en la lista")
                siguiente += 1
                continue
            code, nombre = por_numero[siguiente]
            etiqueta = f"{siguiente:04d}"
            siguiente += 1
        elif re.fullmatch(r"\d{1,4}", base) and int(base) in por_numero:
            code, nombre = por_numero[int(base)]
        else:
            print(f"SALTADA {src}: el nombre '{base}' no es un codigo de barras de la lista ni un N° de foto (usa --desde)")
            continue
        for ext in EXT:  # si ya habia una foto de este producto, se reemplaza
            viejo = os.path.join(DESTINO, code + ext)
            if os.path.exists(viejo):
                os.remove(viejo)
        procesar(src, os.path.join(DESTINO, code + ".webp"))
        print(f"{etiqueta}  {code}  {nombre}  <- {os.path.basename(src)}")


if __name__ == "__main__":
    main()
